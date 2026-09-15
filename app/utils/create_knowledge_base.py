"""
DocMind — Synthetic Knowledge Base Generator

Creates a realistic fictional knowledge base for development,
testing, retrieval experiments, and evaluation.

All information is synthetic and must not be treated as real
company policy or confidential information.
"""

from pathlib import Path

from docx import Document
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
)


PROJECT_ROOT = Path(__file__).resolve().parents[2]
KNOWLEDGE_BASE = PROJECT_ROOT / "knowledge_base"


# ============================================================
# Helper Functions
# ============================================================

def ensure_directories() -> None:
    """Create all knowledge-base directories."""
    for category in ["company", "products", "technical", "hr"]:
        (KNOWLEDGE_BASE / category).mkdir(parents=True, exist_ok=True)


def create_pdf(
    path: Path,
    title: str,
    sections: list[tuple[str, str]],
) -> None:
    """Create a simple text-based PDF document."""
    styles = getSampleStyleSheet()

    document = SimpleDocTemplate(
        str(path),
        pagesize=LETTER,
        rightMargin=50,
        leftMargin=50,
        topMargin=50,
        bottomMargin=50,
    )

    story = []

    story.append(
        Paragraph(title, styles["Title"])
    )
    story.append(Spacer(1, 16))

    for heading, body in sections:
        story.append(
            Paragraph(heading, styles["Heading2"])
        )
        story.append(Spacer(1, 6))

        for paragraph in body.split("\n\n"):
            story.append(
                Paragraph(
                    paragraph.replace("\n", "<br/>"),
                    styles["BodyText"],
                )
            )
            story.append(Spacer(1, 10))

    document.build(story)


def create_docx(
    path: Path,
    title: str,
    sections: list[tuple[str, str]],
) -> None:
    """Create a DOCX document."""
    document = Document()

    document.add_heading(title, level=0)

    for heading, body in sections:
        document.add_heading(heading, level=1)

        for paragraph in body.split("\n\n"):
            document.add_paragraph(paragraph)

    document.save(path)


def create_txt(
    path: Path,
    title: str,
    sections: list[tuple[str, str]],
) -> None:
    """Create a plain-text document."""
    lines = [
        title,
        "=" * len(title),
        "",
    ]

    for heading, body in sections:
        lines.extend(
            [
                heading,
                "-" * len(heading),
                body,
                "",
            ]
        )

    path.write_text(
        "\n".join(lines),
        encoding="utf-8",
    )


# ============================================================
# Knowledge Base Content
# ============================================================

def generate_company_documents() -> None:
    """Generate company-related documents."""

    create_pdf(
        KNOWLEDGE_BASE / "company" / "company_overview.pdf",
        "NovaTech Solutions — Company Overview",
        [
            (
                "Company Profile",
                (
                    "NovaTech Solutions is a fictional technology company "
                    "created for the DocMind RAG project. The company "
                    "provides software development, cloud integration, "
                    "data engineering, and technical consulting services."
                ),
            ),
            (
                "Mission",
                (
                    "NovaTech Solutions aims to help organizations build "
                    "reliable, maintainable, and secure software systems "
                    "through practical engineering and responsible use "
                    "of technology."
                ),
            ),
            (
                "Departments",
                (
                    "The company is organized into Engineering, Product, "
                    "Customer Success, Human Resources, Finance, and "
                    "Operations departments."
                ),
            ),
            (
                "Working Model",
                (
                    "NovaTech uses a hybrid working model. Teams may work "
                    "remotely or from designated company offices according "
                    "to their role and manager-approved schedule."
                ),
            ),
        ],
    )

    create_docx(
        KNOWLEDGE_BASE / "company" / "services.docx",
        "NovaTech Solutions — Services",
        [
            (
                "Software Development",
                (
                    "NovaTech provides custom web application, mobile "
                    "application, API, and backend development services. "
                    "Projects normally begin with requirements gathering "
                    "and continue through development, testing, deployment, "
                    "and maintenance."
                ),
            ),
            (
                "Cloud Integration",
                (
                    "The company provides cloud architecture and integration "
                    "services for organizations migrating applications to "
                    "managed cloud environments."
                ),
            ),
            (
                "Data Engineering",
                (
                    "Data engineering services include data pipelines, "
                    "ETL workflows, database integration, validation, and "
                    "analytics infrastructure."
                ),
            ),
            (
                "Technical Consulting",
                (
                    "Technical consulting engagements can include "
                    "architecture reviews, performance analysis, security "
                    "reviews, and engineering process improvements."
                ),
            ),
            (
                "Service Support",
                (
                    "Customer support requests are categorized by severity. "
                    "Critical production incidents receive the highest "
                    "priority and are escalated to the technical response "
                    "team."
                ),
            ),
        ],
    )

    create_pdf(
        KNOWLEDGE_BASE / "company" / "policies.pdf",
        "NovaTech Solutions — General Company Policies",
        [
            (
                "Information Security",
                (
                    "Employees must protect company information and must "
                    "not share confidential credentials, private customer "
                    "data, or internal documents with unauthorized persons."
                ),
            ),
            (
                "Acceptable Technology Use",
                (
                    "Company systems should primarily be used for legitimate "
                    "business activities. Employees must not intentionally "
                    "install malicious software or bypass security controls."
                ),
            ),
            (
                "Remote Work",
                (
                    "Employees working remotely are responsible for "
                    "maintaining a secure working environment. Company "
                    "systems should be accessed through approved devices "
                    "and security controls."
                ),
            ),
            (
                "Policy Questions",
                (
                    "Questions about company-wide policies should be "
                    "directed to Human Resources or the relevant department "
                    "owner."
                ),
            ),
        ],
    )


def generate_product_documents() -> None:
    """Generate product-related documents."""

    create_pdf(
        KNOWLEDGE_BASE / "products" / "product_a.pdf",
        "NovaTech Product A — NovaDesk",
        [
            (
                "Overview",
                (
                    "NovaDesk is a fictional customer-support management "
                    "platform. It allows support teams to manage customer "
                    "requests, assign tickets, track statuses, and monitor "
                    "service performance."
                ),
            ),
            (
                "Key Features",
                (
                    "NovaDesk includes ticket management, team assignment, "
                    "priority levels, customer profiles, dashboards, "
                    "search, and configurable notification rules."
                ),
            ),
            (
                "User Roles",
                (
                    "NovaDesk provides Administrator, Manager, Agent, "
                    "and Viewer roles. Administrators configure the system, "
                    "while Agents primarily work on assigned customer tickets."
                ),
            ),
            (
                "Default Ticket Priorities",
                (
                    "Tickets can be classified as Low, Medium, High, or "
                    "Critical. Critical tickets represent severe production "
                    "issues requiring immediate attention."
                ),
            ),
        ],
    )

    create_txt(
        KNOWLEDGE_BASE / "products" / "product_b.txt",
        "NovaTech Product B — NovaFlow",
        [
            (
                "Overview",
                (
                    "NovaFlow is a fictional workflow automation platform "
                    "designed to help organizations automate repetitive "
                    "business processes."
                ),
            ),
            (
                "Workflow Structure",
                (
                    "A NovaFlow workflow consists of a trigger, one or more "
                    "actions, optional conditions, and a final execution "
                    "state."
                ),
            ),
            (
                "Supported Actions",
                (
                    "Common actions include sending notifications, updating "
                    "records, calling HTTP endpoints, creating tasks, and "
                    "writing approved data to external systems."
                ),
            ),
            (
                "Execution States",
                (
                    "A workflow can be Pending, Running, Completed, Failed, "
                    "or Disabled."
                ),
            ),
        ],
    )

    create_txt(
        KNOWLEDGE_BASE / "products" / "faq.txt",
        "NovaTech Products — Frequently Asked Questions",
        [
            (
                "NovaDesk Access",
                (
                    "NovaDesk administrators can invite users from the "
                    "Administration section. The invited user receives "
                    "instructions for completing account setup."
                ),
            ),
            (
                "NovaFlow Failure",
                (
                    "When a NovaFlow workflow fails, administrators should "
                    "review the execution log, identify the failed action, "
                    "and correct the configuration before retrying."
                ),
            ),
            (
                "Priority Levels",
                (
                    "NovaDesk supports Low, Medium, High, and Critical "
                    "ticket priorities."
                ),
            ),
            (
                "Product Documentation",
                (
                    "Product-specific configuration instructions should be "
                    "followed before changing production settings."
                ),
            ),
        ],
    )


def generate_technical_documents() -> None:
    """Generate technical documentation."""

    create_pdf(
        KNOWLEDGE_BASE / "technical" / "installation.pdf",
        "NovaTech Technical Guide — Installation",
        [
            (
                "System Requirements",
                (
                    "NovaDesk requires a supported modern web browser and "
                    "network access to the NovaTech application environment. "
                    "Administrative configuration requires an administrator "
                    "account."
                ),
            ),
            (
                "Installation Process",
                (
                    "The recommended installation process is to verify "
                    "system requirements, configure the environment, "
                    "initialize the application database, create an "
                    "administrator account, and perform a health check."
                ),
            ),
            (
                "Health Check",
                (
                    "After installation, administrators should verify that "
                    "the application is reachable, the database connection "
                    "is operational, and background services are running."
                ),
            ),
        ],
    )

    create_docx(
        KNOWLEDGE_BASE / "technical" / "configuration.docx",
        "NovaTech Technical Guide — Configuration",
        [
            (
                "Application Configuration",
                (
                    "NovaDesk configuration is divided into application, "
                    "database, authentication, notification, and logging "
                    "settings."
                ),
            ),
            (
                "Authentication",
                (
                    "Authentication settings control how users sign in. "
                    "Administrative changes to authentication configuration "
                    "should be tested before deployment to production."
                ),
            ),
            (
                "Notifications",
                (
                    "Notification rules determine when users receive "
                    "messages about ticket assignments, status changes, "
                    "and critical incidents."
                ),
            ),
            (
                "Logging",
                (
                    "Application logs should contain enough information "
                    "for administrators to diagnose failures without "
                    "recording passwords, authentication tokens, or other "
                    "secrets."
                ),
            ),
        ],
    )

    create_txt(
        KNOWLEDGE_BASE / "technical" / "troubleshooting.txt",
        "NovaTech Technical Guide — Troubleshooting",
        [
            (
                "Application Does Not Start",
                (
                    "Check the application configuration, verify that "
                    "required services are available, inspect application "
                    "logs, and confirm that the database connection is "
                    "working."
                ),
            ),
            (
                "Database Connection Failure",
                (
                    "Verify the database address, credentials, network "
                    "connectivity, and database service status. Never "
                    "include database passwords in support tickets."
                ),
            ),
            (
                "Users Cannot Sign In",
                (
                    "Check authentication configuration, account status, "
                    "identity-provider connectivity, and relevant logs."
                ),
            ),
            (
                "Notifications Are Missing",
                (
                    "Verify notification rules, recipient configuration, "
                    "background processing, and application logs."
                ),
            ),
        ],
    )


def generate_hr_documents() -> None:
    """Generate HR-related documents."""

    create_pdf(
        KNOWLEDGE_BASE / "hr" / "leave_policy.pdf",
        "NovaTech Solutions — Employee Leave Policy",
        [
            (
                "Annual Leave",
                (
                    "Full-time employees receive 20 annual leave days per "
                    "calendar year. Annual leave is intended for planned "
                    "personal time away from work."
                ),
            ),
            (
                "Sick Leave",
                (
                    "Employees receive 10 sick leave days per calendar year. "
                    "Sick leave should be used when an employee is unable "
                    "to work because of illness."
                ),
            ),
            (
                "Casual Leave",
                (
                    "Employees receive 5 casual leave days per calendar "
                    "year for short-term personal matters that do not "
                    "require planned annual leave."
                ),
            ),
            (
                "Request Procedure",
                (
                    "Planned annual leave should normally be submitted "
                    "through the HR portal at least 5 working days before "
                    "the requested start date."
                ),
            ),
            (
                "Approval",
                (
                    "Leave requests are reviewed by the employee's direct "
                    "manager. Approval depends on staffing requirements "
                    "and operational needs."
                ),
            ),
            (
                "Carryover",
                (
                    "Unused annual leave may be carried into the next "
                    "calendar year subject to manager approval and applicable "
                    "company rules."
                ),
            ),
        ],
    )

    create_docx(
        KNOWLEDGE_BASE / "hr" / "attendance_policy.docx",
        "NovaTech Solutions — Attendance Policy",
        [
            (
                "Working Hours",
                (
                    "Standard full-time employees are expected to work "
                    "8 hours per working day, excluding meal breaks."
                ),
            ),
            (
                "Attendance Recording",
                (
                    "Employees should record attendance using the company's "
                    "approved attendance system."
                ),
            ),
            (
                "Late Arrival",
                (
                    "Employees who expect to arrive late should notify "
                    "their manager as soon as reasonably possible."
                ),
            ),
            (
                "Remote Attendance",
                (
                    "Employees working remotely must remain available "
                    "during their agreed working schedule and maintain "
                    "reliable communication with their team."
                ),
            ),
            (
                "Unplanned Absence",
                (
                    "Employees unable to attend work unexpectedly should "
                    "notify their manager and follow the applicable leave "
                    "procedure."
                ),
            ),
        ],
    )

    create_pdf(
        KNOWLEDGE_BASE / "hr" / "employee_handbook.pdf",
        "NovaTech Solutions — Employee Handbook",
        [
            (
                "Employment Categories",
                (
                    "NovaTech uses full-time, part-time, contract, and "
                    "intern employment categories. Benefits and policies "
                    "may vary according to employment category."
                ),
            ),
            (
                "Leave Summary",
                (
                    "Full-time employees receive 20 annual leave days, "
                    "10 sick leave days, and 5 casual leave days per "
                    "calendar year, subject to the detailed leave policy."
                ),
            ),
            (
                "Leave Requests",
                (
                    "Planned leave should normally be requested through "
                    "the HR portal. Employees should provide sufficient "
                    "notice so managers can plan team coverage."
                ),
            ),
            (
                "Manager Responsibility",
                (
                    "Managers are responsible for reviewing leave requests "
                    "and balancing employee needs with operational staffing "
                    "requirements."
                ),
            ),
            (
                "Professional Conduct",
                (
                    "Employees are expected to communicate respectfully, "
                    "protect company information, and follow applicable "
                    "company policies."
                ),
            ),
        ],
    )


# ============================================================
# Main
# ============================================================

def main() -> None:
    """Generate the complete synthetic knowledge base."""
    print("=" * 70)
    print("DOCMIND — SYNTHETIC KNOWLEDGE BASE GENERATOR")
    print("=" * 70)

    ensure_directories()

    generate_company_documents()
    generate_product_documents()
    generate_technical_documents()
    generate_hr_documents()

    files = sorted(KNOWLEDGE_BASE.rglob("*"))

    documents = [
        file
        for file in files
        if file.is_file()
    ]

    print()
    print(f"Knowledge base created at:")
    print(f"  {KNOWLEDGE_BASE}")
    print()
    print(f"Documents generated: {len(documents)}")
    print()

    for document in documents:
        print(f"  ✓ {document.relative_to(KNOWLEDGE_BASE)}")

    print()
    print("=" * 70)
    print("KNOWLEDGE BASE GENERATION COMPLETE")
    print("=" * 70)


if __name__ == "__main__":
    main()